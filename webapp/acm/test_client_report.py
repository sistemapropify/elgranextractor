import io
import zipfile
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from docx import Document
from PIL import Image

from .client_report import prepare_report, author_name
from .components_engine import calculate, candidates
from .components_pdf import build_acm_pdf
from .components_report import build_acm_docx
from .components_html import build_acm_html
from .test_components import params, sample, _pdf_text


class ClientReportTests(SimpleTestCase):
    def setUp(self):
        self.p = {**params(), 'property_type': 'Casa', 'manual_valuation': 234000,
                  'address': 'Calle de prueba 123', 'report_notes': 'Buena conservación.'}
        self.rows = candidates(sample(), self.p)
        self.result = calculate(self.rows, self.p)
        self.user = SimpleNamespace(first_name='María Alejandra', last_name='López Gómez', username='cuenta-interna')

    def word_text(self, content):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return archive.read('word/document.xml').decode()

    def test_full_name_is_used_in_both_reports_without_copying_the_reference_author(self):
        self.assertEqual(author_name(self.user), 'María Alejandra López Gómez')
        pdf = _pdf_text(build_acm_pdf(self.p, self.rows, self.result, user=self.user, fetch_images=False))
        word = self.word_text(build_acm_docx(self.p, self.rows, self.result, user=self.user, fetch_images=False))
        html = build_acm_html(self.p, self.rows, self.result, user=self.user)
        self.assertIn(b'Preparado por:', pdf)
        self.assertIn(b'Alejandra', pdf)
        self.assertIn('María Alejandra López Gómez', word)
        for text in (word, html):
            self.assertIn('María Alejandra López Gómez', text)
        for text in (pdf.decode('latin1'), word, html):
            self.assertIn('Agente Inmobiliario', text)
            self.assertNotIn('cuenta-interna', text)
            self.assertNotIn('Erika Palacios', text)

    def test_both_formats_preserve_the_old_three_final_amounts(self):
        pdf = _pdf_text(build_acm_pdf(self.p, self.rows, self.result, fetch_images=False)).decode('latin1')
        word = self.word_text(build_acm_docx(self.p, self.rows, self.result, fetch_images=False))
        for text in (pdf, word):
            for amount in ('245,700', '234,000', '222,300'):
                self.assertIn(amount, text)
            self.assertLess(text.index('VALOR DE SALIDA AL MERCADO'), text.index('VALOR COMERCIAL'))
            self.assertNotIn('Nuevo precio', text)

    def test_word_is_portrait_and_contains_no_source_metadata_or_internal_tables(self):
        for row in self.rows:
            row.update(source='remax', code='SECRET-CODE', url='https://remax.pe/secret',
                       district='Cayma Remax', first_seen='2020-01-01')
        body = build_acm_docx(self.p, self.rows, self.result, fetch_images=False)
        text = self.word_text(body)
        for hidden in ('remax', 'SECRET-CODE', '2020-01-01', 'Ajuste terreno', 'Fuentes consultadas'):
            self.assertNotIn(hidden, text)
        doc = Document(io.BytesIO(body))
        self.assertLess(doc.sections[0].page_width, doc.sections[0].page_height)

    def test_only_selected_comparable_photos_are_loaded_and_never_a_subject_cadastral_photo(self):
        image = io.BytesIO(); Image.new('RGB', (8, 6), 'teal').save(image, 'PNG')
        self.p.update(report_ids=[self.rows[0]['id']], subject_lot_id='04012607037012',
                      _acm_permissions={'report_photos': True, 'photos': True})
        for row in self.rows: row['image'] = 'https://approved.example/'+row['id']+'.jpg'
        with patch('acm.components_pdf._property_photo', return_value=image.getvalue()) as photo, \
             patch('acm.components_pdf._static_map', return_value=image.getvalue()):
            data = prepare_report(self.p, self.rows, self.result)
        photo.assert_called_once()
        self.assertEqual(photo.call_args.args[0], self.rows[0]['image'])
        self.assertNotIn('subject', data)
        self.assertEqual(len(data['photos']), 1)

    def test_photo_permission_applies_equally_to_pdf_and_word(self):
        image = io.BytesIO(); Image.new('RGB', (8, 6), 'teal').save(image, 'PNG')
        for row in self.rows: row['image'] = 'https://approved.example/photo.jpg'
        with patch('acm.components_pdf._property_photo', return_value=image.getvalue()) as photo, \
             patch('acm.components_pdf._static_map', return_value=image.getvalue()):
            for builder in (build_acm_pdf, build_acm_docx):
                for permitted in (False, True):
                    photo.reset_mock()
                    self.p['_acm_permissions'] = {'report_photos': permitted, 'photos': True}
                    builder(self.p, self.rows, self.result)
                    self.assertEqual(photo.called, permitted)

    def test_date_uses_peru_time_and_absent_data_is_not_invented(self):
        data = prepare_report(self.p, [], {'new': None}, generated_at=datetime(2026, 10, 17, 2, tzinfo=timezone.utc), fetch_images=False)
        self.assertEqual(data['date'].day, 16)
        self.assertIsNone(data['values'])
        self.assertNotIn('expectation', data)

    def test_long_observations_are_preserved_in_both_formats(self):
        self.p['report_notes'] = '\n'.join('Observación '+str(i)+': ubicación y conservación revisadas.' for i in range(100))
        pdf = _pdf_text(build_acm_pdf(self.p, self.rows, self.result, fetch_images=False))
        word = self.word_text(build_acm_docx(self.p, self.rows, self.result, fetch_images=False))
        self.assertIn(b'99:', pdf)
        self.assertIn('99:', word)

    def test_one_comparable_keeps_five_columns_and_a_small_photo(self):
        self.p['report_ids'] = [self.rows[0]['id']]
        image = io.BytesIO(); Image.new('RGB', (1200, 800), 'teal').save(image, 'PNG')
        self.p['_acm_permissions'] = {'report_photos': True, 'photos': True}
        self.rows[0]['image'] = 'https://approved.example/photo.jpg'
        with patch('acm.components_pdf._property_photo', return_value=image.getvalue()), \
             patch('acm.components_pdf._static_map', return_value=image.getvalue()):
            body = build_acm_docx(self.p, self.rows, self.result)
        doc = Document(io.BytesIO(body))
        cards = next(t for t in doc.tables if len(t.columns) == 5)
        self.assertEqual(len(cards.rows), 1)
        self.assertLess(cards.rows[0].cells[0].width.inches, 1.5)
        # The comparable photograph is the last image after logo, map and chart.
        self.assertLess(doc.inline_shapes[-1].width.inches, 1.5)
        self.assertLessEqual(doc.inline_shapes[-1].height.inches, .64)

    def test_subject_type_uses_one_photo_strip_without_an_extra_land_strip(self):
        data = prepare_report(self.p, self.rows, self.result, fetch_images=False)
        self.assertEqual(len(data['groups']), 1)
        self.assertTrue(all(row['kind'] == 'Casa' for row in data['rows']))
        self.assertLessEqual(len(data['rows']), 5)

    def test_pdm_use_zone_and_parameter_tables_are_in_pdf_word_and_html(self):
        pdm = {
            'codigo': 'RDA-2',
            'nombre': 'Residencial densidad alta tipo 2',
            'categoria': 'Residencial',
            'color': '#d6882d',
            'property_type': 'Casa',
            'confidence': 'alta',
            'approximate': False,
            'distance_edge': 20,
            'typologies': ['Multifamiliar'],
            'rows': [
                {'label': 'Densidad', 'values': ['De 2251 a 2800 hab/ha']},
                {'label': 'Lote mínimo', 'values': ['600.00 m2']},
                {'label': 'Altura', 'values': ['10 pisos']},
            ],
            'compatible_uses': 'CZ, CM, E-1, H1, H2, ZR',
            'notes': ['Ficha normativa PDM.'],
        }
        with patch('acm.client_report.pdm_zone', return_value=pdm), \
             patch('acm.components_html.pdm_zone', return_value=pdm):
            pdf = _pdf_text(build_acm_pdf(
                self.p, self.rows, self.result, fetch_images=False))
            word = self.word_text(build_acm_docx(
                self.p, self.rows, self.result, fetch_images=False))
            html = build_acm_html(self.p, self.rows, self.result)
        outputs = (('pdf', pdf.decode('latin1')), ('word', word), ('html', html))
        for kind, text in outputs:
            for value in ('Zona de uso PDM', 'RDA-2', 'Densidad',
                          'Usos compatibles', 'Multifamiliar'):
                self.assertIn(value, text)
            self.assertIn('PROPIEDADES COMPARABLES' if kind == 'pdf'
                          else 'Propiedades comparables', text)
            self.assertLess(text.index('Justificaci'),
                            text.index('Zona de uso PDM'))
            self.assertLess(text.index('Zona de uso PDM'),
                            text.index('Agente Inmobiliario'))
            if kind == 'html':
                self.assertLess(text.index('Valor de salida al mercado'),
                                text.index('Inmueble a valorar'))
            else:
                valuation = text.index('Valorizaci')
                self.assertLess(valuation, text.index('Ubicaci', valuation))
        for text in (word, html):
            self.assertIn('Lote mínimo', text)
