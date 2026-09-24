"""Informe Word del ACM por componentes con las operaciones que sustentan el resultado."""
from datetime import datetime
from io import BytesIO
from statistics import median


def _number(value, decimals=0):
    if value is None:
        return 'Sin dato'
    return f'{float(value):,.{decimals}f}'


def _money(value):
    return 'Sin dato' if value is None else f'USD {float(value):,.0f}'


def _unit(value):
    return 'Sin dato' if value is None else f'USD {float(value):,.0f}/m²'


def _median_text(values):
    values = sorted(float(value) for value in values)
    if not values:
        return 'No existen valores aptos.'
    ordered = ', '.join(_unit(value) for value in values)
    if len(values) % 2:
        explanation = f'El valor central es {_unit(values[len(values)//2])}.'
    else:
        left, right = values[len(values)//2-1:len(values)//2+1]
        explanation = f'Los dos valores centrales son {_unit(left)} y {_unit(right)}; su promedio es {_unit((left+right)/2)}.'
    return f'Valores ordenados: {ordered}. {explanation}'


def build_acm_docx(params, records, result, excluded=(), generated_at=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    generated_at = generated_at or datetime.now()
    excluded = set(excluded)
    by_id = {row['id']: row for row in records}
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(.55); section.bottom_margin = Inches(.55)
    section.left_margin = Inches(.55); section.right_margin = Inches(.55)

    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[style_name]
        style.font.name = 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(9.5)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(12)

    def shade(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        node = tc_pr.find(qn('w:shd')) or OxmlElement('w:shd')
        node.set(qn('w:fill'), fill)
        if node.getparent() is None: tc_pr.append(node)

    def borders(cell):
        tc_pr = cell._tc.get_or_add_tcPr()
        edges = tc_pr.first_child_found_in('w:tcBorders')
        if edges is None:
            edges = OxmlElement('w:tcBorders'); tc_pr.append(edges)
        for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge = OxmlElement(f'w:{name}')
            edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9')
            edges.append(edge)

    def table(headers, rows, widths=None):
        tbl = doc.add_table(rows=1, cols=len(headers))
        tbl.autofit = True
        for index, text in enumerate(headers):
            cell = tbl.rows[0].cells[index]; cell.text = str(text); shade(cell, '244A68'); borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(8)
            if widths: cell.width = Inches(widths[index])
        for row_index, values in enumerate(rows):
            cells = tbl.add_row().cells
            for index, text in enumerate(values):
                cells[index].text = str(text); borders(cells[index]); cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index % 2: shade(cells[index], 'F2F6FA')
                for paragraph in cells[index].paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs: run.font.size = Pt(8)
                if widths: cells[index].width = Inches(widths[index])
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return tbl

    def heading(text, level=1):
        paragraph = doc.add_heading(text, level=level)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)
        return paragraph

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run('Informe de análisis comparativo de mercado')
    intro = doc.add_paragraph()
    intro.add_run('Propósito. ').bold = True
    intro.add_run('Explicar con los datos seleccionados cómo se obtuvo la estimación, qué anuncios participaron y cuáles fueron las operaciones realizadas. Los importes corresponden a precios de oferta.')
    doc.add_paragraph(f'Generado el {generated_at:%d/%m/%Y %H:%M}. Método: {result.get("version", "componentes")}')

    heading('Inmueble objetivo')
    target_rows = [
        ('Tipo', params['property_type']), ('Radio de propiedades', f'{_number(params["radius"])} m'),
        ('Área de terreno', f'{_number(params.get("land"),2)} m²' if params.get('land') else 'No aplica'),
        ('Área construida', f'{_number(params.get("built"),2)} m²' if params.get('built') else 'No aplica'),
        ('Fuentes consultadas', ', '.join(params.get('sources', []))),
    ]
    table(('Dato', 'Valor'), target_rows, (2.2, 4.7))

    heading('Resultado')
    new = result.get('new')
    if new:
        if result['model'] == 'components':
            result_rows = [('Valor del terreno objetivo', _money(new['land_value'])),
                           ('Valor de construcción y mejoras', _money(new['built_value'])),
                           ('Estimación total', _money(new['total'])),
                           ('Rango central orientativo', f'{_money(new["range_low"])} a {_money(new["range_high"])}')]
        else:
            result_rows = [('Valor unitario usado', _unit(new['unit'])),
                           ('Estimación total', _money(new['total'])),
                           ('Rango central orientativo', f'{_money(new["range_low"])} a {_money(new["range_high"])}')]
        table(('Concepto', 'Resultado'), result_rows, (3.8, 3.1))
    else:
        doc.add_paragraph('No se obtuvo una estimación final porque faltaron comparables aptos para completar el método.')

    if result['model'] == 'components':
        heading('Paso 1 Cálculo del valor del suelo')
        lands = [by_id[row_id] for row_id in result.get('land_ids', []) if row_id in by_id and row_id not in excluded]
        doc.add_paragraph('Para cada terreno usado se divide el precio anunciado entre su área de terreno. Después se ordenan esos valores y se toma la mediana. La mediana es el valor central y evita que un anuncio extremadamente caro o barato domine el resultado.')
        land_rows = []
        land_units = []
        for row in lands:
            unit = row['price']/row['land']; land_units.append(unit)
            land_rows.append((row['source'].upper(), row.get('code') or row['id'], f'{_number(row["distance"])} m',
                              _money(row['price']), f'{_number(row["land"],2)} m²', _unit(unit)))
        if land_rows:
            table(('Portal', 'Código', 'Distancia', 'Precio', 'Terreno', 'Precio por m²'), land_rows, (.7,1.1,.8,1.15,1,1.2))
            doc.add_paragraph(_median_text(land_units))
            paragraph = doc.add_paragraph(); paragraph.add_run('Valor de suelo adoptado: ').bold = True; paragraph.add_run(_unit(result['land_unit']))
        else:
            doc.add_paragraph('No hubo terrenos aptos. Sin una referencia de suelo el método no puede separar el valor del terreno del valor de la construcción.')

        heading('Paso 2 Valor que cada casa atribuye a construcción y mejoras')
        doc.add_paragraph('A cada casa se le descuenta el valor estimado de su terreno. El dinero restante se divide entre el área construida. Ese remanente incluye construcción, antigüedad, estado, distribución, piscina y otras mejoras; no es un costo de obra certificado.')
        breakdown = {row['id']: row for row in result.get('breakdown', [])}
        house_rows = []
        built_units = []
        for row_id in result.get('house_ids', []):
            if row_id not in by_id or row_id not in breakdown or row_id in excluded: continue
            row, detail = by_id[row_id], breakdown[row_id]
            if detail['usable']: built_units.append(detail['built_unit'])
            house_rows.append((row['source'].upper(), row.get('code') or row['id'], _money(row['price']),
                               f'{_number(row["land"],2)} m²', _money(detail['land_value']),
                               _money(detail['remainder']), f'{_number(row["built"],2)} m²',
                               f'{_number(detail.get("land_similarity"),1)}%', f'{_number(detail.get("built_similarity"),1)}%',
                               _unit(detail['built_unit']), _money(detail['target_estimate']),
                               'Sí' if detail['usable'] else 'No'))
        if house_rows:
            table(('Portal','Código','Precio','Terreno','Valor suelo','Remanente','Construido','Sim. terreno','Sim. construcción','Remanente por m²','Valor sugerido','Usada'), house_rows,
                  (.5,.7,.75,.6,.75,.75,.65,.65,.75,.9,.8,.4))
            if built_units:
                if result.get('built_unit_method') == 'closest_comparable':
                    chosen_id = result.get('built_reference_id')
                    chosen = by_id.get(chosen_id, {})
                    chosen_detail = breakdown.get(chosen_id, {})
                    doc.add_paragraph(
                        f'Solo hubo dos casas y sus aportes quedaron muy dispersos. Se eligió el comparable más parecido '
                        f'por superficies y distancia ({chosen.get("source", "").upper()} · {chosen.get("code", chosen_id)}), '
                        f'en lugar de promediar los dos extremos. El otro anuncio queda como referencia.'
                    )
                    paragraph = doc.add_paragraph(); paragraph.add_run('Aporte unitario adoptado para construcción y mejoras: ').bold = True; paragraph.add_run(_unit(chosen_detail.get('built_unit')))
                elif result.get('built_unit_method') == 'weighted_median':
                    chosen_id = result.get('built_reference_id')
                    chosen = by_id.get(chosen_id, {})
                    chosen_detail = breakdown.get(chosen_id, {})
                    doc.add_paragraph('Los aportes están muy dispersos. Se usó una mediana ponderada por similitud: las casas con terreno, construcción y ubicación más parecidos tienen mayor peso y no se promedian los extremos.')
                    paragraph = doc.add_paragraph(); paragraph.add_run('Aporte unitario adoptado para construcción y mejoras: ').bold = True; paragraph.add_run(_unit(chosen_detail.get('built_unit')))
                else:
                    doc.add_paragraph(_median_text(built_units))
                    paragraph = doc.add_paragraph(); paragraph.add_run('Aporte unitario adoptado para construcción y mejoras: ').bold = True; paragraph.add_run(_unit(median(built_units)))
            else:
                doc.add_paragraph('Ninguna casa dejó un remanente positivo después de descontar el suelo. Por eso no se estimó el aporte de construcción y mejoras.')
        else:
            doc.add_paragraph('No hubo casas completas y aptas dentro del radio seleccionado.')

        heading('Paso 3 Aplicación al inmueble objetivo')
        if new:
            doc.add_paragraph(f'Terreno: {_number(params["land"],2)} m² × {_unit(result["land_unit"])} = {_money(new["land_value"])}.')
            doc.add_paragraph(f'Construcción y mejoras: {_number(params["built"],2)} m² × {_unit(new["built_unit"])} = {_money(new["built_value"])}.')
            final = doc.add_paragraph(); final.add_run('Estimación total: ').bold = True
            final.add_run(f'{_money(new["land_value"])} + {_money(new["built_value"])} = {_money(new["total"])}.')
        else:
            doc.add_paragraph('El proceso se detuvo antes de este paso porque no existió evidencia suficiente para uno de los componentes.')
    else:
        heading('Proceso de cálculo')
        area_key = 'land' if result['model'] == 'land' else 'built'
        label = 'terreno' if area_key == 'land' else 'área construida'
        selected_ids = result.get('land_ids', []) or result.get('house_ids', [])
        selected = [by_id[row_id] for row_id in selected_ids if row_id in by_id and row_id not in excluded]
        rows = [(row['source'].upper(), row.get('code') or row['id'], _money(row['price']),
                 f'{_number(row[area_key],2)} m²', _unit(row['price']/row[area_key])) for row in selected]
        doc.add_paragraph(f'Cada anuncio se convierte a precio por m² de {label}. Se ordenan los valores y se toma la mediana. Finalmente, la mediana se multiplica por el área objetivo.')
        if rows: table(('Portal','Código','Precio','Área','Precio por m²'), rows, (.8,1.2,1.4,1.2,1.5))
        if new: doc.add_paragraph(f'{_number(params[area_key],2)} m² × {_unit(new["unit"])} = {_money(new["total"])}.')

    heading('Registros que no participaron')
    reference = [row for row in records if row['id'] not in set(result.get('land_ids', [])) | set(result.get('house_ids', [])) or row['id'] in excluded]
    reasons = {}
    for row in reference:
        row_reasons = ['Desmarcada manualmente'] if row['id'] in excluded else (row.get('issues') or ['Fuera del grupo finalmente utilizado'])
        for reason in row_reasons: reasons[reason] = reasons.get(reason, 0) + 1
    stored_summary = result.get('reference_summary') or {}
    if stored_summary:
        reasons = stored_summary.get('reasons', reasons)
        reference_count = stored_summary.get('count', len(reference))
    else:
        reference_count = len(reference)
    doc.add_paragraph(f'{reference_count} registros quedaron visibles como referencia y no modificaron el resultado.')
    if reasons:
        table(('Cantidad', 'Motivo'), [(count, reason) for reason, count in sorted(reasons.items(), key=lambda item: -item[1])], (1,5.9))

    heading('Alcance de la estimación')
    doc.add_paragraph('Este cálculo usa precios publicados, no precios finales de compraventa. Solo participan ubicaciones exactas y registros con precio y superficies necesarias. Habitaciones, baños y piso se muestran como información, pero no excluyen comparables ni agregan valor mientras no exista un ajuste económico validado. Una muestra pequeña se calcula y se señala para revisión.')
    if result.get('messages'):
        heading('Advertencias del análisis', level=2)
        for message in result['messages']: doc.add_paragraph(message, style='List Bullet')

    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()
