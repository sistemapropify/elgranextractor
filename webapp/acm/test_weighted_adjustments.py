import json
from io import BytesIO
from unittest.mock import patch
from django.test import SimpleTestCase
from docx import Document
from acm.components_engine import candidates,calculate
from acm.components_views import recalculate,search,word_report
from acm.components_report import build_acm_docx
from acm.test_components import record,params,ComponentsEndpointTests


def fixture():
    p={**params(),'property_type':'Casa','land':210,'built':250}
    rows=candidates([record('a',price=200000,land=200,built=230),
                     record('land','Terreno',price=120000,land=200)],p)
    return p,rows


class WeightedAdjustmentTests(SimpleTestCase):
    def test_example_preserves_price_then_adds_only_surface_differences(self):
        p,rows=fixture();r=calculate(rows,p);d=r['breakdown'][0]
        self.assertAlmostEqual(d['land_adjustment'],6000)
        self.assertAlmostEqual(d['built_adjustment'],80000/230*20)
        self.assertAlmostEqual(r['new']['total'],212956.52173913)
        self.assertEqual(d['similarity_weight'],100)

    def test_fixed_weights_increase_and_decrease_monotonically(self):
        p,rows=fixture()
        rows+=candidates([record('b',price=310000,land=280,built=350)],p)
        values=[];weights=[]
        for land,built in [(190,210),(200,230),(210,250),(220,270)]:
            r=calculate(rows,{**p,'weight_reference':p,'land':land,'built':built})
            values.append(r['new']['total'])
            weights.append([d['similarity_weight'] for d in r['breakdown']])
        self.assertEqual(values,sorted(values))
        self.assertTrue(all(a<b for a,b in zip(values,values[1:])))
        self.assertTrue(all(w==weights[0] for w in weights))

    def test_weights_sum_to_100_and_contributions_sum_to_result(self):
        p,rows=fixture()
        rows+=candidates([record('b',price=600000,land=390,built=490)],p)
        r=calculate(rows,p);details=r['breakdown']
        self.assertGreater(details[0]['similarity_weight'],details[1]['similarity_weight'])
        self.assertAlmostEqual(sum(d['similarity_weight'] for d in details),100)
        self.assertAlmostEqual(sum(d['weighted_contribution'] for d in details),r['new']['total'])
        self.assertGreaterEqual(r['new']['total'],min(d['target_estimate'] for d in details))
        self.assertLessEqual(r['new']['total'],max(d['target_estimate'] for d in details))

    def test_word_contains_actual_adjustments_and_weights(self):
        p,rows=fixture();r=calculate(rows,p)
        doc=Document(BytesIO(build_acm_docx(p,rows,r)))
        text=' '.join(p.text for p in doc.paragraphs)+' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn('Casa usada',text)
        self.assertIn('Ajuste terreno',text)
        self.assertNotIn('Las demás casas quedan como referencia',text)


class ScenarioEndpointTests(ComponentsEndpointTests):
    @patch('acm.components_views._persist_history')
    @patch('acm.components_report.build_acm_docx',return_value=b'word')
    @patch('acm.components_views.load_records')
    def test_scenario_and_word_use_same_signed_evidence(self,load,build,persist):
        from types import SimpleNamespace
        p,rows=fixture();load.return_value=(rows,[])
        persist.return_value=(SimpleNamespace(codigo_display='ACM-test'),True)
        data=json.loads(search(self.request(p)).content)
        payload={'token':data['token'],'target_areas':{'land':220,'built':270}}
        changed=json.loads(recalculate(self.request(payload)).content)
        self.assertGreater(changed['result']['new']['total'],data['result']['new']['total'])
        self.assertEqual(changed['params']['weight_reference']['land'],210)
        word_report(self.request(payload))
        self.assertEqual(build.call_args.args[2],changed['result'])
        self.assertEqual(build.call_args.args[0],changed['params'])

    @patch('acm.components_views.load_records')
    def test_cannot_supply_weights_or_invalid_areas(self,load):
        p,rows=fixture();load.return_value=(rows,[])
        data=json.loads(search(self.request(p)).content)
        self.assertEqual(recalculate(self.request({'token':data['token'],'target_areas':{'land':-1,'built':200}})).status_code,400)
